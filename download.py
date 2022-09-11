import pytube
from pytube import YouTube
path = executable_path=r"C:\Users\KDK\Desktop\ytbe_dnwload"
import pandas as pd
from openpyxl import load_workbook

y=2
while y < 409:
    df = pd.read_excel(r"C:\Users\KDK\Desktop\ytbe_dnwload\links.xlsx")
    wb = load_workbook(r"C:\Users\KDK\Desktop\ytbe_dnwload\links.xlsx")
    ws = wb.worksheets[0]
    links = ws.cell(row=y, column=3).value
    print(links)
   ## links2 =
    y+=1
    yt=YouTube(links)
    t=yt.streams.filter(only_audio=True).all()
    t[0].download(links)
    print(y)

yt=YouTube("https://www.youtube.com/watch?v=UncaiC2XDbQ")
t=yt.streams.filter(only_audio=True).all()
t[0].download("https://www.youtube.com/watch?v=UncaiC2XDbQ")
